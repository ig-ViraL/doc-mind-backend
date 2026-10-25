"""
Document parser — extracts per-page text (with tables, image descriptions and
links folded in) plus a list of headings for building the section tree.

PDF:  page text via PyMuPDF, tables via pdfplumber, headings from the PDF outline
      (bookmarks) or, when there is none, from font sizes.
DOCX: "Heading N" paragraph styles; text is split into virtual pages.
TXT:  markdown / numbered / ALL-CAPS heading heuristics; virtual pages.

Blocking libraries run in a worker thread; vision LLM calls run concurrently
on the event loop, capped by MAX_VISION_IMAGES and VISION_CONCURRENCY.
"""

from __future__ import annotations

import asyncio
import base64
import io
import logging
import re
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

from config import settings

logger = logging.getLogger(__name__)

VIRTUAL_PAGE_CHARS = 3000


@dataclass
class Heading:
    level: int
    title: str
    page: int  # 1-based page the heading starts on


@dataclass
class ParsedDocument:
    filename: str
    file_type: str
    pages: list[str]                       # index 0 = page 1
    headings: list[Heading] = field(default_factory=list)
    has_real_pages: bool = True            # False for DOCX/TXT virtual pages
    links: list[str] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def text(self) -> str:
        return "\n\n".join(p for p in self.pages if p.strip())


async def parse_bytes(raw_bytes: bytes, filename: str, content_type: str = "") -> ParsedDocument:
    """Parse raw file bytes. Raises ValueError for oversized, unsupported or broken files."""
    max_bytes = settings.max_file_size_mb * 1024 * 1024
    if len(raw_bytes) > max_bytes:
        raise ValueError(
            f"File '{filename}' exceeds the {settings.max_file_size_mb} MB limit "
            f"({len(raw_bytes) // (1024 * 1024)} MB)."
        )

    name = filename.lower()
    logger.info("Parsing '%s' (content_type=%s, size=%d bytes)", filename, content_type, len(raw_bytes))

    if content_type == "application/pdf" or name.endswith(".pdf"):
        return await _parse_pdf(raw_bytes, filename)
    if (
        content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or name.endswith(".docx")
    ):
        return await _parse_docx(raw_bytes, filename)
    if content_type.startswith("text/plain") or name.endswith((".txt", ".md")):
        return _parse_txt(raw_bytes, filename)

    logger.warning("Unsupported content_type '%s' — attempting unstructured fallback", content_type)
    return await asyncio.to_thread(_parse_unstructured, raw_bytes, filename, content_type)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _sync_extract_pdf(raw_bytes: bytes) -> dict[str, Any]:
    import fitz  # pymupdf
    import pdfplumber

    doc = fitz.open(stream=raw_bytes, filetype="pdf")
    try:
        pages_text: list[str] = []
        links: list[str] = []
        raw_images: list[tuple[int, bytes]] = []
        for pno, page in enumerate(doc, start=1):
            pages_text.append(page.get_text("text").strip())
            for link in page.get_links():
                if link.get("uri"):
                    links.append(link["uri"])
            for img_info in page.get_images(full=True):
                try:
                    img_bytes = doc.extract_image(img_info[0])["image"]
                    if len(img_bytes) >= settings.min_image_bytes:
                        raw_images.append((pno, img_bytes))
                except Exception:
                    pass

        headings = [
            Heading(level=max(1, int(lvl)), title=str(title).strip(), page=max(1, int(pg)))
            for lvl, title, pg in doc.get_toc(simple=True)
            if str(title).strip() and int(pg) >= 1
        ]
        if len(headings) < 2:
            headings = _font_headings(doc)
    finally:
        doc.close()

    tables_by_page: dict[int, list[str]] = {}
    try:
        with pdfplumber.open(io.BytesIO(raw_bytes)) as pdf:
            for pno, page in enumerate(pdf.pages, start=1):
                md = [_table_to_markdown(t) for t in (page.extract_tables() or []) if t]
                if md:
                    tables_by_page[pno] = md
    except Exception as exc:
        logger.warning("PDF table extraction skipped: %s", exc)

    return {
        "pages_text": pages_text,
        "links": links,
        "raw_images": raw_images,
        "headings": headings,
        "tables_by_page": tables_by_page,
    }


_NUMBERED_RE = re.compile(r"^(\d+(?:\.\d+){0,3})\.?\s+\S")


def _font_headings(doc: Any) -> list[Heading]:
    """Detect headings from font size / weight when the PDF has no outline."""
    size_chars: Counter[float] = Counter()
    lines: list[tuple[int, str, float, bool]] = []
    for pno, page in enumerate(doc, start=1):
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = [s for s in line.get("spans", []) if s.get("text", "").strip()]
                if not spans:
                    continue
                text = " ".join(s["text"].strip() for s in spans).strip()
                size = round(max(s["size"] for s in spans), 1)
                bold = all(s.get("flags", 0) & 16 for s in spans)
                for s in spans:
                    size_chars[round(s["size"], 1)] += len(s["text"])
                lines.append((pno, text, size, bold))
    if not size_chars:
        return []

    body = size_chars.most_common(1)[0][0]
    page_total = max(1, len(doc))
    # Running headers/footers repeat on many pages; they are not headings.
    repeats = Counter(t for _, t, _, _ in lines)

    candidates: list[tuple[int, str, float]] = []
    for pno, text, size, bold in lines:
        if not (3 <= len(text) <= 120) or len(text.split()) > 14:
            continue
        if text.endswith((".", ",", ";")) or not any(c.isalpha() for c in text):
            continue
        if repeats[text] > max(2, page_total * 0.3):
            continue
        big = size >= body * 1.15
        bold_numbered = bold and size >= body and _NUMBERED_RE.match(text)
        if big or bold_numbered:
            candidates.append((pno, text, size))

    if not candidates or len(candidates) > page_total * 8:
        return []

    sizes = sorted({s for _, _, s in candidates}, reverse=True)
    level_of = {s: min(i + 1, 3) for i, s in enumerate(sizes)}
    headings: list[Heading] = []
    for pno, text, size in candidates:
        level = level_of[size]
        m = _NUMBERED_RE.match(text)
        if m:
            level = min(m.group(1).count(".") + 1, 4)
        # Merge a title that wraps onto a second line.
        if headings and headings[-1].page == pno and headings[-1].level == level and len(headings[-1].title) < 60 and not m:
            headings[-1].title += " " + text
            continue
        headings.append(Heading(level=level, title=text, page=pno))
    return headings


async def _parse_pdf(raw_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        data = await asyncio.to_thread(_sync_extract_pdf, raw_bytes)
    except Exception as exc:
        logger.exception("PDF parsing failed for '%s': %s", filename, exc)
        raise ValueError(f"Failed to parse PDF '{filename}': {exc}") from exc

    images_by_page = await _describe_images(data["raw_images"])

    pages: list[str] = []
    for pno, page_text in enumerate(data["pages_text"], start=1):
        parts = [page_text] if page_text else []
        for i, md in enumerate(data["tables_by_page"].get(pno, []), start=1):
            parts.append(f"[TABLE {i}]\n{md}")
        for i, desc in enumerate(images_by_page.get(pno, []), start=1):
            parts.append(f"[IMAGE {i}]\n{desc}")
        pages.append("\n\n".join(parts))

    links = list(dict.fromkeys(data["links"]))
    if links and pages:
        pages[-1] += "\n\n[HYPERLINKS]\n" + "\n".join(f"• {u}" for u in links)

    logger.info(
        "PDF '%s': pages=%d headings=%d tables=%d images=%d links=%d",
        filename, len(pages), len(data["headings"]),
        sum(len(v) for v in data["tables_by_page"].values()),
        sum(len(v) for v in images_by_page.values()), len(links),
    )
    return ParsedDocument(
        filename=filename, file_type="pdf", pages=pages, headings=data["headings"], links=links
    )


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _sync_extract_docx(raw_bytes: bytes) -> dict[str, Any]:
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(io.BytesIO(raw_bytes))

    blocks: list[tuple[str, int | None]] = []  # (text, heading level or None)
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        level: int | None = None
        if style == "Title":
            level = 1
        else:
            m = re.match(r"Heading (\d)", style)
            if m:
                level = int(m.group(1))
        blocks.append((text, level))

    for i, table in enumerate(doc.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _table_to_markdown(rows)
        if md:
            blocks.append((f"[TABLE {i}]\n{md}", None))

    raw_images: list[tuple[int, bytes]] = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype.lower():
            try:
                blob = rel.target_part.blob
                if len(blob) >= settings.min_image_bytes:
                    raw_images.append((0, blob))
            except Exception:
                pass

    links: list[str] = []
    try:
        for para in doc.paragraphs:
            for hl in para._element.findall(f'.//{{{qn("w")}}}hyperlink'):
                r_id = hl.get(f'{{{qn("r")}}}id')
                if r_id and r_id in doc.part.rels:
                    links.append(doc.part.rels[r_id].target_ref)
    except Exception as exc:
        logger.warning("DOCX link extraction skipped: %s", exc)

    return {"blocks": blocks, "raw_images": raw_images, "links": links}


async def _parse_docx(raw_bytes: bytes, filename: str) -> ParsedDocument:
    try:
        data = await asyncio.to_thread(_sync_extract_docx, raw_bytes)
    except Exception as exc:
        logger.exception("DOCX parsing failed for '%s': %s", filename, exc)
        raise ValueError(f"Failed to parse DOCX '{filename}': {exc}") from exc

    blocks: list[tuple[str, int | None]] = list(data["blocks"])
    images = (await _describe_images(data["raw_images"])).get(0, [])
    blocks += [(f"[IMAGE {i}]\n{d}", None) for i, d in enumerate(images, start=1)]
    links = list(dict.fromkeys(data["links"]))
    if links:
        blocks.append(("[HYPERLINKS]\n" + "\n".join(f"• {u}" for u in links), None))

    if not any(level for _, level in blocks):
        # No styled headings: fall back to the plain-text heuristics.
        detected = _text_heading_levels([t for t, _ in blocks])
        blocks = [(t, detected.get(i)) for i, (t, _) in enumerate(blocks)]

    pages, headings = _virtual_pages(blocks)
    return ParsedDocument(
        filename=filename, file_type="docx", pages=pages, headings=headings,
        has_real_pages=False, links=links,
    )


# ---------------------------------------------------------------------------
# TXT
# ---------------------------------------------------------------------------

_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*$")


def _text_heading_levels(paragraphs: list[str]) -> dict[int, int]:
    """Map paragraph index -> heading level for single-line paragraphs that look like headings."""
    levels: dict[int, int] = {}
    for i, para in enumerate(paragraphs):
        line = para.strip()
        if "\n" in line or not (2 <= len(line) <= 100):
            continue
        m = _MD_HEADING_RE.match(line)
        if m:
            levels[i] = len(m.group(1))
            continue
        if line.endswith((".", ",", ";", ":")):
            continue
        m = _NUMBERED_RE.match(line)
        if m and len(line.split()) <= 12 and line[len(m.group(1)):].strip()[:1].isupper():
            levels[i] = min(m.group(1).count(".") + 1, 4)
            continue
        letters = [c for c in line if c.isalpha()]
        if len(letters) >= 4 and all(c.isupper() for c in letters) and len(line.split()) <= 10:
            levels[i] = 1
    return levels


def _virtual_pages(blocks: list[tuple[str, int | None]]) -> tuple[list[str], list[Heading]]:
    """Pack blocks into ~VIRTUAL_PAGE_CHARS pages, recording which page each heading lands on."""
    pages: list[str] = []
    headings: list[Heading] = []
    current: list[str] = []
    size = 0
    for text, level in blocks:
        if current and size + len(text) > VIRTUAL_PAGE_CHARS:
            pages.append("\n\n".join(current))
            current, size = [], 0
        if level:
            title = _MD_HEADING_RE.sub(r"\2", text).strip()
            headings.append(Heading(level=level, title=title, page=len(pages) + 1))
        current.append(text)
        size += len(text) + 2
    if current:
        pages.append("\n\n".join(current))
    return pages, headings


def _parse_txt(raw_bytes: bytes, filename: str) -> ParsedDocument:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            text = raw_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Unable to decode text file '{filename}' with any supported encoding")

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    # A heading line glued to its paragraph ("# Title\nbody") becomes its own block.
    split: list[str] = []
    for p in paragraphs:
        first, _, rest = p.partition("\n")
        if rest and _text_heading_levels([first]):
            split += [first, rest]
        else:
            split.append(p)
    levels = _text_heading_levels(split)
    pages, headings = _virtual_pages([(p, levels.get(i)) for i, p in enumerate(split)])
    return ParsedDocument(
        filename=filename, file_type="txt", pages=pages or [""], headings=headings, has_real_pages=False
    )


# ---------------------------------------------------------------------------
# Unstructured fallback
# ---------------------------------------------------------------------------


def _parse_unstructured(raw_bytes: bytes, filename: str, content_type: str) -> ParsedDocument:
    try:
        from unstructured.partition.auto import partition

        elements = partition(file=io.BytesIO(raw_bytes), content_type=content_type)
        blocks: list[tuple[str, int | None]] = []
        for el in elements:
            text = str(el).strip()
            if text:
                blocks.append((text, 1 if type(el).__name__ == "Title" else None))
        pages, headings = _virtual_pages(blocks)
        return ParsedDocument(
            filename=filename, file_type=content_type.split("/")[-1] or "unknown",
            pages=pages or [""], headings=headings, has_real_pages=False,
        )
    except Exception as exc:
        logger.exception("Unstructured parsing failed for '%s': %s", filename, exc)
        raise ValueError(
            f"Unsupported file type '{content_type}' for '{filename}'. Supported types: PDF, DOCX, TXT."
        ) from exc


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _table_to_markdown(rows: list[list[Any]]) -> str:
    """Convert a list-of-rows into a GitHub-flavoured Markdown table."""
    if not rows:
        return ""

    def clean(cell: Any) -> str:
        return str(cell or "").replace("|", "\\|").replace("\n", " ").strip()

    cleaned = [[clean(cell) for cell in row] for row in rows]
    col_count = max(len(row) for row in cleaned)
    padded = [row + [""] * (col_count - len(row)) for row in cleaned]
    header, body = padded[0], padded[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
        *("| " + " | ".join(row) + " |" for row in body),
    ]
    return "\n".join(lines)


async def _describe_images(raw_images: list[tuple[int, bytes]]) -> dict[int, list[str]]:
    """Describe up to MAX_VISION_IMAGES images (largest first), VISION_CONCURRENCY at a time."""
    if not settings.extract_images or not raw_images or not settings.openrouter_api_key:
        return {}
    chosen = sorted(raw_images, key=lambda x: len(x[1]), reverse=True)[: settings.max_vision_images]
    chosen.sort(key=lambda x: x[0])
    if len(raw_images) > len(chosen):
        logger.info("Describing %d of %d images (MAX_VISION_IMAGES)", len(chosen), len(raw_images))
    sem = asyncio.Semaphore(settings.vision_concurrency)

    async def one(img: bytes) -> str:
        async with sem:
            return await _describe_image(img)

    results = await asyncio.gather(*[one(img) for _, img in chosen], return_exceptions=True)
    out: dict[int, list[str]] = {}
    for (page, _), res in zip(chosen, results):
        if isinstance(res, str) and res:
            out.setdefault(page, []).append(res)
    return out


async def _describe_image(img_bytes: bytes) -> str:
    """Send one image to the vision LLM; returns "" on failure."""
    from services.llm_client import get_openrouter_client

    try:
        if img_bytes[:8] == b"\x89PNG\r\n\x1a\n":
            mime = "image/png"
        elif img_bytes[:6] in (b"GIF87a", b"GIF89a"):
            mime = "image/gif"
        elif img_bytes[8:12] == b"WEBP":
            mime = "image/webp"
        else:
            mime = "image/jpeg"

        b64 = base64.b64encode(img_bytes).decode("utf-8")
        response = await get_openrouter_client().chat.completions.create(
            model=settings.vision_model,
            max_tokens=300,
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                        {
                            "type": "text",
                            "text": (
                                "Describe this image concisely. "
                                "If it contains text, transcribe it exactly. "
                                "If it is a chart or graph, summarize the data and trends. "
                                "If it is a diagram or flowchart, explain the structure."
                            ),
                        },
                    ],
                }
            ],
        )
        return (response.choices[0].message.content or "").strip()
    except Exception as exc:
        logger.warning("Image description failed: %s", exc)
        return ""
